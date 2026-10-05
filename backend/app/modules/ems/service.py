"""Event registry use-cases: creation from a goal (GQM), versioning with approval, lifecycle, import, inference."""

from __future__ import annotations

import csv
import io
import json
import uuid
from collections import Counter, defaultdict
from datetime import UTC, datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ConflictError, NotFoundError, PermissionDenied, ValidationFailed
from app.modules.audit.service import record
from app.modules.ems import semver
from app.modules.ems.models import Event, EventVersion, GlobalParam, TrackingConfig
from app.modules.iam.models import Project
from app.modules.iam.permissions import P
from app.modules.iam.policy import Principal

TRANSITIONS = {
    "active": {"deprecated"},
    "deprecated": {"active", "archived"},
    "draft": {"archived"},
    "archived": set(),
}


async def get_event(
    db: AsyncSession, principal: Principal, event_id: uuid.UUID, permission: str = P.EVENTS_VIEW
) -> Event:
    ev = await db.get(Event, event_id)
    if ev is None or ev.org_id != principal.org_id:
        raise NotFoundError("Событие не найдено")
    visible = principal.visible_project_ids()
    if visible is not None and ev.project_id not in visible:
        raise NotFoundError("Событие не найдено")
    principal.require(permission, ev.project_id)
    return ev


async def versions_of(db: AsyncSession, event_id: uuid.UUID) -> list[EventVersion]:
    q = (
        select(EventVersion)
        .where(EventVersion.event_id == event_id)
        .order_by(EventVersion.major.desc(), EventVersion.minor.desc(), EventVersion.patch.desc())
    )
    return list((await db.execute(q)).scalars())


def latest_approved(versions: list[EventVersion]) -> EventVersion | None:
    return next((v for v in versions if v.status == "approved"), None)


async def create_event(
    db: AsyncSession,
    principal: Principal,
    project: Project,
    *,
    name: str,
    description: str = "",
    category: str = "",
    owner: str = "",
    goal: str = "",
    question: str = "",
    metric_keys: list[str] | None = None,
    tags: list[str] | None = None,
    params: list[dict[str, Any]] | None = None,
    app_version: str = "",
) -> tuple[Event, EventVersion]:
    principal.require(P.EVENTS_EDIT, project.id)
    if not semver.NAME_RE.match(name):
        raise ValidationFailed("Имя события должно быть в snake_case (a-z, 0-9, _), например match_start")
    if await db.scalar(select(Event.id).where(Event.project_id == project.id, Event.name == name)):
        raise ConflictError(f"Событие {name} уже есть в реестре проекта")
    clean = semver.normalize_params(params or [])
    ev = Event(
        org_id=principal.org_id,
        project_id=project.id,
        name=name,
        description=description,
        category=category,
        owner=owner or principal.label,
        goal=goal,
        question=question,
        metric_keys=metric_keys or [],
        tags=tags or [],
        status="draft",
        created_by=principal.label,
    )
    db.add(ev)
    await db.flush()
    v = EventVersion(
        event_id=ev.id,
        version="1.0.0",
        major=1,
        minor=0,
        patch=0,
        params=clean,
        description=description,
        changelog="Первая версия",
        app_version=app_version,
        author=principal.label,
    )
    db.add(v)
    await db.flush()
    await record(
        db,
        "event.create",
        principal=principal,
        resource_type="event",
        resource_id=ev.id,
        project_id=project.id,
        details={"name": name},
    )
    return ev, v


async def propose_version(
    db: AsyncSession,
    principal: Principal,
    ev: Event,
    params: list[dict[str, Any]],
    *,
    description: str | None = None,
    changelog: str = "",
    app_version: str = "",
) -> EventVersion:
    principal.require(P.EVENTS_EDIT, ev.project_id)
    if ev.status == "archived":
        raise ConflictError("Событие в архиве")
    versions = await versions_of(db, ev.id)
    if any(v.status == "pending" for v in versions):
        raise ConflictError("Есть неутверждённая версия — дождитесь решения или отклоните её")
    base = latest_approved(versions)
    clean = semver.normalize_params(params)
    d = semver.diff(base.params if base else [], clean)
    desc = description if description is not None else (base.description if base else ev.description)
    if d.bump == "none" and base and desc == base.description:
        raise ValidationFailed("Изменений нет")
    version = semver.next_version(base.version if base else None, d.bump if d.bump != "none" else "patch")
    major, minor, patch = semver.parse(version)
    v = EventVersion(
        event_id=ev.id,
        version=version,
        major=major,
        minor=minor,
        patch=patch,
        params=clean,
        description=desc,
        changelog=changelog,
        app_version=app_version,
        author=principal.label,
    )
    db.add(v)
    await db.flush()
    await record(
        db,
        "event.version_proposed",
        principal=principal,
        resource_type="event",
        resource_id=ev.id,
        project_id=ev.project_id,
        details={"version": version, **d.as_dict()},
    )
    return v


async def review(
    db: AsyncSession, principal: Principal, ev: Event, version_id: uuid.UUID, approve: bool, comment: str
) -> EventVersion:
    principal.require(P.EVENTS_APPROVE, ev.project_id)
    v = await db.get(EventVersion, version_id)
    if v is None or v.event_id != ev.id:
        raise NotFoundError("Версия не найдена")
    if v.status != "pending":
        raise ConflictError("Версия уже рассмотрена")
    v.status = "approved" if approve else "rejected"
    v.reviewed_by, v.reviewed_at, v.review_comment = principal.label, datetime.now(UTC), comment
    if approve:
        ev.current_version = v.version
        ev.description = v.description or ev.description
        if ev.status == "draft":
            ev.status = "active"
    await record(
        db,
        "event.version_reviewed",
        principal=principal,
        resource_type="event",
        resource_id=ev.id,
        project_id=ev.project_id,
        details={"version": v.version, "approved": approve, "comment": comment},
    )
    return v


async def change_status(db: AsyncSession, principal: Principal, ev: Event, status: str, reason: str) -> Event:
    principal.require(P.EVENTS_APPROVE, ev.project_id)
    if status not in TRANSITIONS.get(ev.status, set()):
        raise ConflictError(f"Переход {ev.status} → {status} недопустим")
    before = ev.status
    ev.status = status
    await record(
        db,
        "event.status_changed",
        principal=principal,
        resource_type="event",
        resource_id=ev.id,
        project_id=ev.project_id,
        details={"from": before, "to": status, "reason": reason},
    )
    return ev


async def global_params(db: AsyncSession, org_id: uuid.UUID) -> list[GlobalParam]:
    return list(
        (await db.execute(select(GlobalParam).where(GlobalParam.org_id == org_id).order_by(GlobalParam.name))).scalars()
    )


def export_view(ev: Event, v: EventVersion, project_key: str) -> dict[str, Any]:
    return {
        "name": ev.name,
        "description": v.description or ev.description,
        "version": v.version,
        "project": project_key,
        "params": v.params,
    }


# ------------------------------------------------------------------ import of existing tracking plans
COLUMNS = {
    "event_name": ("event_name", "event", "событие", "имя события"),
    "event_description": ("event_description", "описание события"),
    "param_name": ("param_name", "param", "параметр"),
    "param_type": ("param_type", "type", "тип"),
    "required": ("required", "param_required", "обязательный"),
    "param_description": ("param_description", "описание параметра", "description"),
    "enum": ("enum", "values", "значения"),
}
TYPE_ALIASES = {
    "str": "string",
    "text": "string",
    "integer": "int",
    "number": "float",
    "double": "float",
    "boolean": "bool",
    "datetime": "timestamp",
    "date": "timestamp",
    "object": "json",
}


def _rows_from_file(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return []
        rows = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = content.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    if not rows:
        return []
    header = [h.lower() for h in rows[0]]
    mapping: dict[str, int] = {}
    for key, aliases in COLUMNS.items():
        for i, h in enumerate(header):
            if h in aliases:
                mapping[key] = i
                break
    if "event_name" not in mapping:
        raise ValidationFailed("В файле нет колонки event_name (или «событие»)")
    return [{k: (r[i] if i < len(r) else "") for k, i in mapping.items()} for r in rows[1:] if any(r)]


def parse_tracking_plan(filename: str, content: bytes) -> dict[str, dict[str, Any]]:
    events: dict[str, dict[str, Any]] = {}
    for r in _rows_from_file(filename, content):
        name = r["event_name"].strip().lower()
        if not name:
            continue
        ev = events.setdefault(name, {"description": "", "params": []})
        if r.get("event_description"):
            ev["description"] = r["event_description"]
        pname = r.get("param_name", "").strip().lower()
        if pname:
            ptype = TYPE_ALIASES.get(
                r.get("param_type", "string").lower(), r.get("param_type", "string").lower() or "string"
            )
            enum = [x.strip() for x in r.get("enum", "").replace(";", ",").split(",") if x.strip()]
            ev["params"].append(
                {
                    "name": pname,
                    "type": "enum" if enum and ptype == "string" else ptype,
                    "required": r.get("required", "").lower() in ("1", "true", "yes", "да", "+", "required"),
                    "description": r.get("param_description", ""),
                    "enum": enum,
                }
            )
    return events


async def import_plan(
    db: AsyncSession, principal: Principal, project: Project, plan: dict[str, dict[str, Any]], app_version: str = ""
) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {"created": [], "versioned": [], "unchanged": [], "errors": []}
    existing = {e.name: e for e in (await db.execute(select(Event).where(Event.project_id == project.id))).scalars()}
    for name, spec in plan.items():
        try:
            if name in existing:
                await propose_version(
                    db,
                    principal,
                    existing[name],
                    spec["params"],
                    description=spec["description"] or None,
                    changelog="Импорт разметки",
                    app_version=app_version,
                )
                result["versioned"].append(name)
            else:
                await create_event(
                    db,
                    principal,
                    project,
                    name=name,
                    description=spec["description"],
                    params=spec["params"],
                    app_version=app_version,
                )
                result["created"].append(name)
        except ValidationFailed as exc:
            (result["unchanged"] if exc.message == "Изменений нет" else result["errors"]).append(
                f"{name}: {exc.message}"
            )
        except ConflictError as exc:
            result["errors"].append(f"{name}: {exc.message}")
    return result


# ------------------------------------------------------------------ inference from actual data
def infer_params(samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Types and required-ness from observed JSON params (sample of real events)."""
    n = len(samples)
    seen: dict[str, Counter[str]] = defaultdict(Counter)
    values: dict[str, set[str]] = defaultdict(set)
    for s in samples:
        for k, v in s.items():
            if v is None:
                continue
            t = (
                "bool"
                if isinstance(v, bool)
                else "int"
                if isinstance(v, int)
                else "float"
                if isinstance(v, float)
                else "json"
                if isinstance(v, (dict, list))
                else "string"
            )
            seen[k][t] += 1
            if t == "string" and len(values[k]) <= 20:
                values[k].add(str(v))
    out = []
    for k, types in sorted(seen.items()):
        t = "float" if set(types) == {"int", "float"} else types.most_common(1)[0][0]
        enum = sorted(values[k]) if t == "string" and n >= 30 and 0 < len(values[k]) <= 8 else []
        out.append(
            {
                "name": k,
                "type": "enum" if enum else t,
                "required": sum(types.values()) >= 0.99 * n,
                "description": "",
                "enum": enum,
            }
        )
    return out


def parse_params(raw: Any) -> dict[str, Any]:
    if isinstance(raw, dict):
        return raw
    if not raw:
        return {}
    try:
        val = json.loads(raw)
    except (TypeError, ValueError):
        return {}
    return val if isinstance(val, dict) else {}


async def tracking_config(db: AsyncSession, project_id: uuid.UUID) -> TrackingConfig | None:
    return await db.get(TrackingConfig, project_id)


def require_permission(principal: Principal, permission: str, project_id: uuid.UUID) -> None:
    if not principal.can(permission, project_id):
        raise PermissionDenied("Недостаточно прав", details={"permission": permission})


def quoted(cfg: TrackingConfig, attr: str) -> str:
    """Column name from the tracking config as a quoted identifier (config is admin-provided, still sanitised)."""
    return '"' + str(getattr(cfg, attr)).replace('"', "") + '"'
