"""Plain text from documents people actually write release notes in: Markdown/TXT, Word, PDF."""

from __future__ import annotations

import io
from pathlib import Path

from app.core.errors import ValidationFailed

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 60_000  # ~15k tokens: a long release document still fits the model context
SUPPORTED = (".txt", ".md", ".markdown", ".docx", ".pdf")


def extract_text(filename: str, content: bytes) -> tuple[str, bool]:
    """Returns (text, truncated)."""
    if len(content) > MAX_BYTES:
        raise ValidationFailed("Файл больше 10 МБ")
    ext = Path(filename.lower()).suffix
    if ext in (".txt", ".md", ".markdown"):
        text = _decode(content)
    elif ext == ".docx":
        text = _docx(content)
    elif ext == ".pdf":
        text = _pdf(content)
    else:
        raise ValidationFailed(
            f"Формат {ext or 'без расширения'} не поддерживается", details={"allowed": list(SUPPORTED)}
        )
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ValidationFailed("В документе не найден текст (для PDF нужен текстовый слой, не скан)")
    return (text[:MAX_CHARS], True) if len(text) > MAX_CHARS else (text, False)


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _docx(content: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValidationFailed("Не удалось прочитать .docx") from exc
    lines: list[str] = []
    for p in doc.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        prefix = "#" * int(style[-1]) + " " if style.startswith("heading") and style[-1:].isdigit() else ""
        if p.text.strip():
            lines.append(prefix + p.text)
    for table in doc.tables:  # requirement tables are common in specs
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ValidationFailed("Не удалось прочитать PDF") from exc
