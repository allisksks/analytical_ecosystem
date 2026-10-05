from __future__ import annotations

import io
import json
from collections.abc import Callable, Iterator
from typing import Any

import pytest
from httpx import AsyncClient

from app.modules.ai.gateway import set_gateway
from app.modules.ems.ai_drafts import clean_params, snake
from app.modules.ems.models import GlobalParam
from tests.ai.fake import FakeLLM
from tests.conftest import World

H = Callable[[str], dict[str, str]]

PLAN = {
    "summary": "В 1.9.0 появляется магазин офферов и вкладки магазина.",
    "events": [
        {
            "action": "create",
            "name": "Offer Purchase",
            "description": "Покупка персонального оффера",
            "category": "monetisation",
            "goal": "Монетизация",
            "question": "Покупают ли офферы?",
            "rationale": "Новая фича",
            "quote": "Персональный оффер на главном экране",
            "params": [
                {"name": "user_id", "type": "string"},
                {"name": "offer_id", "type": "string", "required": True},
                {"name": "price", "type": "number"},
                {"name": "tier", "type": "enum"},
            ],
        },
        {
            "action": "update",
            "name": "shop_open",
            "description": "Открытие магазина",
            "quote": "вкладки магазина",
            "params": [{"name": "tab", "type": "enum", "enum": ["offers", "gems"]}],
        },
        {"action": "update", "name": "shop_open", "params": []},
        {"action": "create", "name": "123", "params": []},
    ],
}


@pytest.fixture
def llm() -> Iterator[FakeLLM]:
    calls = {"n": 0}

    def reply(messages: list[dict[str, str]]) -> str:
        calls["n"] += 1
        if "НЕ JSON" in messages[1]["content"] and calls["n"] == 1:
            return "Конечно! Вот разметка."
        return "```json\n" + json.dumps(PLAN, ensure_ascii=False) + "\n```"

    fake = FakeLLM(reply)
    set_gateway(fake.gateway())
    yield fake
    set_gateway(None)


@pytest.fixture
async def setup(client: AsyncClient, world: World, as_user: H) -> dict[str, Any]:
    world.db.add(GlobalParam(org_id=world.org.id, name="user_id", type="string"))
    await world.add_user("lead", "analytics_lead", None)
    await world.add_user("dev", "developer", ["alpha"])
    await world.db.commit()
    alpha = str(world.projects["alpha"].id)
    r = await client.post(
        "/ems/events",
        json={"project_id": alpha, "name": "shop_open", "params": [{"name": "source", "type": "string"}]},
        headers=as_user("analyst"),
    )
    ev = r.json()
    await client.post(
        f"/ems/events/{ev['id']}/versions/{ev['versions'][0]['id']}/review",
        json={"approve": True},
        headers=as_user("lead"),
    )
    return {"alpha": alpha, "shop_open": ev["id"]}


def test_helpers() -> None:
    assert snake("Offer Purchase") == "offer_purchase" and snake("levelComplete") == "level_complete"
    assert snake("123 go") == "e_123_go" and snake("123") == "" and snake("покупка") == ""
    warnings: list[str] = []
    params = clean_params(
        [{"name": "user_id"}, {"name": "Price", "type": "double"}, {"name": "x", "type": "weird"}],
        {"user_id"},
        warnings,
    )
    assert params == [
        {"name": "price", "type": "float", "required": False, "description": "", "enum": []},
        {"name": "x", "type": "string", "required": False, "description": "", "enum": []},
    ]
    assert len(warnings) == 1


async def test_document_to_reviewed_drafts(
    client: AsyncClient, as_user: H, llm: FakeLLM, setup: dict[str, Any]
) -> None:
    hdr = as_user("analyst")
    r = await client.post(
        "/ems/ai-drafts",
        data={
            "project_id": setup["alpha"],
            "app_version": "1.9.0",
            "text": "Релиз 1.9.0. Персональный оффер на главном экране, вкладки магазина.",
        },
        headers=hdr,
    )
    assert r.status_code == 201, r.text
    d = r.json()
    prompt = llm.bodies[-1]["messages"][1]["content"]
    assert "shop_open" in prompt and "user_id" in prompt and "Персональный оффер" in prompt
    assert d["summary"].startswith("В 1.9.0") and d["pending"] == 2
    create, update = d["items"]
    assert (create["action"], create["name"]) == ("create", "offer_purchase")
    assert [p["name"] for p in create["params"]] == ["offer_id", "price", "tier"]  # global user_id dropped
    assert any("tier" in w for w in create["warnings"])  # enum without values
    assert (update["action"], update["event_id"]) == ("update", setup["shop_open"])
    assert [p["name"] for p in update["params"]] == ["source", "tab"] and update["diff"]["bump"] == "minor"

    # the analyst edits, then accepts: registry drafts with pending versions appear for the lead
    r = await client.patch(f"/ems/ai-drafts/{d['id']}/items/{create['id']}", json={"name": "offer_buy"}, headers=hdr)
    assert r.json()["items"][0]["name"] == "offer_buy"
    r = await client.post(f"/ems/ai-drafts/{d['id']}/items/{create['id']}/accept", headers=hdr)
    item = r.json()["items"][0]
    assert item["status"] == "accepted" and item["result_version"] == "1.0.0"
    ev = (await client.get(f"/ems/events/{item['event_id']}", headers=hdr)).json()
    assert ev["status"] == "draft" and ev["pending_version"] == "1.0.0" and "ai-draft" in ev["tags"]
    assert ev["versions"][0]["app_version"] == "1.9.0" and "Релиз 1.9.0" in ev["versions"][0]["changelog"]
    r = await client.post(f"/ems/ai-drafts/{d['id']}/items/{update['id']}/accept", headers=hdr)
    assert r.json()["items"][1]["result_version"] == "1.1.0"
    assert (await client.post(f"/ems/ai-drafts/{d['id']}/items/{update['id']}/reject", headers=hdr)).status_code == 409

    lst = (await client.get("/ems/ai-drafts", params={"project_id": setup["alpha"]}, headers=hdr)).json()
    assert lst[0]["accepted"] == 2 and lst[0]["pending"] == 0


async def test_non_json_answer_is_repaired_and_rights(
    client: AsyncClient, as_user: H, llm: FakeLLM, setup: dict[str, Any]
) -> None:
    r = await client.post(
        "/ems/ai-drafts", data={"project_id": setup["alpha"], "text": "НЕ JSON сначала"}, headers=as_user("analyst")
    )
    assert r.status_code == 201 and len(r.json()["items"]) == 2
    assert (
        await client.post("/ems/ai-drafts", data={"project_id": setup["alpha"], "text": "x"}, headers=as_user("dev"))
    ).status_code == 403
    assert (
        await client.post("/ems/ai-drafts", data={"project_id": setup["alpha"]}, headers=as_user("analyst"))
    ).status_code == 422


async def test_docx_upload(client: AsyncClient, as_user: H, llm: FakeLLM, setup: dict[str, Any]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Релиз 1.9.0", level=1)
    doc.add_paragraph("Персональный оффер на главном экране.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Фича", "Вкладки магазина"
    buf = io.BytesIO()
    doc.save(buf)
    r = await client.post(
        "/ems/ai-drafts",
        data={"project_id": setup["alpha"]},
        files={"file": ("release.docx", buf.getvalue())},
        headers=as_user("analyst"),
    )
    d = r.json()
    assert r.status_code == 201, r.text
    assert d["filename"] == "release.docx" and d["title"] == "release.docx"
    assert "# Релиз 1.9.0" in d["source_text"] and "Фича | Вкладки магазина" in d["source_text"]
    r = await client.post(
        "/ems/ai-drafts",
        data={"project_id": setup["alpha"]},
        files={"file": ("x.exe", b"MZ")},
        headers=as_user("analyst"),
    )
    assert r.status_code == 422
